"""Generate comparison, report and slides exclusively from experiment artifacts."""

import argparse
import json
from pathlib import Path
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from sklearn.metrics import roc_curve
from pptx import Presentation
from pptx.util import Inches
from .utils import save_json


def report(output="."):
    root = Path(output)
    results = root / "results"
    names = ["resnet50", "densenet121"]
    records = [
        json.loads((results / f"{name}_metrics.json").read_text(encoding="utf-8")) for name in names
    ]
    audit = json.loads((results / "dataset_audit.json").read_text(encoding="utf-8"))
    fields = ["accuracy", "precision", "recall", "specificity", "f1", "auc", "tn", "fp", "fn", "tp"]
    comparison = pd.DataFrame(
        [dict(model=r["model"], **{k: r[k] for k in fields}) for r in records]
    )
    comparison.rename(columns={"recall": "sensitivity"}).to_csv(
        results / "model_comparison.csv", index=False
    )
    best = max(
        records,
        key=lambda r: (
            r["validation"]["auc"],
            r["validation"]["recall"],
            r["validation"]["specificity"],
        ),
    )
    save_json(
        results / "selected_model.json",
        dict(
            model=best["model"],
            criterion="Validation AUC, then sensitivity, then specificity; threshold fixed at 0.5",
        ),
    )
    for name in names:
        frame = pd.read_csv(results / "predictions" / f"{name}.csv")
        fpr, tpr, _ = roc_curve(frame.label, frame.probability)
        plt.plot(fpr, tpr, label=name)
    plt.plot([0, 1], [0, 1], "--", label="Random")
    plt.xlabel("FPR")
    plt.ylabel("TPR")
    plt.legend()
    plt.savefig(results / "roc" / "model_comparison_roc.png", dpi=160)
    plt.close()
    comparison.set_index("model")[["accuracy", "recall", "specificity", "f1", "auc"]].plot.bar(
        rot=0, ylim=(0, 1)
    )
    plt.tight_layout()
    plt.savefig(results / "plots" / "model_comparison.png", dpi=160)
    plt.close()
    lines = [
        "# Phân loại viêm phổi từ X-quang ngực",
        "",
        "Kết quả dưới đây được đọc từ checkpoint, history và dự đoán test của chương trình.",
        "",
        "## Phương pháp",
        "Ảnh grayscale lặp thành 3 kênh, resize 224 × 224, normalize ImageNet. Augmentation nhẹ chỉ trên train. "
        "ResNet50 dùng residual skip connections; DenseNet121 dùng dense connectivity để tái sử dụng đặc trưng. "
        "Cả hai khởi tạo ImageNet, train classifier với backbone và BatchNorm đóng băng, sau đó fine-tune toàn mạng. "
        "AdamW, ReduceLROnPlateau và early stopping theo validation loss; khôi phục checkpoint tốt nhất. "
        "Weighted BCE dùng N_NORMAL/N_PNEUMONIA tính riêng trên train. Positive = PNEUMONIA.",
        "",
        "## Phân bố dữ liệu",
        "```",
        (results / "dataset_counts.csv").read_text(),
        "```",
        f'Số ảnh nguồn: {audit["source_images"]}; giữ lại {audit["retained_images"]}; loại {audit["excluded_images"]}. '
        f'Chi tiết lý do: {audit.get("exclusion_reasons", "Xem dataset_audit.json")}. '
        "Patient ID suy ra từ tên ảnh, không phải metadata bệnh án đã xác minh. "
        "Danh sách ảnh loại: excluded_images.csv. Test gốc được giữ riêng và loại ảnh trùng pixel.",
        "![Phân bố](dataset_distribution.png)",
        "",
        "## Kết quả test",
        "| Model | Accuracy | Precision | Sensitivity | Specificity | F1 | AUC | FN |",
        "|---|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for r in records:
        lines.append(
            "| "
            + r["model"]
            + " | "
            + " | ".join(f"{r[k]:.4f}" for k in fields[:6])
            + f' | {r["fn"]} |'
        )
    lines += [
        "",
        f'Mô hình chọn bằng validation: **{best["model"]}**. Test chỉ dùng đánh giá cuối cùng.',
        "![So sánh](plots/model_comparison.png)",
        "![ROC](roc/model_comparison_roc.png)",
    ]
    first, second = records
    lines += [
        "",
        "### Đối chiếu các tiêu chí",
        f'DenseNet121 − ResNet50 trên test: Accuracy {(second["accuracy"]-first["accuracy"])*100:+.2f} điểm phần trăm; '
        f'Sensitivity {(second["recall"]-first["recall"])*100:+.2f} điểm phần trăm; '
        f'Specificity {(second["specificity"]-first["specificity"])*100:+.2f} điểm phần trăm; '
        f'F1 {second["f1"]-first["f1"]:+.4f}; AUC {second["auc"]-first["auc"]:+.4f}. '
        f'Số FN tương ứng ResNet50/DenseNet121 là {first["fn"]}/{second["fn"]}, '
        f'FP là {first["fp"]}/{second["fp"]}. '
        "Mô hình giảm FN nhưng tăng FP thể hiện đánh đổi giữa sensitivity và specificity; "
        "không kết luận mô hình phù hợp lâm sàng chỉ từ Accuracy hoặc từ một lần chạy.",
    ]
    for r in records:
        name = r["model"]
        history = pd.read_csv(root / "logs" / f"history_{name}.csv")
        head = history[history.stage == 1].sort_values("val_loss").iloc[0]
        fine = history[history.stage == 2].sort_values("val_loss").iloc[0]
        lines += [
            "",
            f"## {name}",
            f'Checkpoint epoch {r["epoch"]}; test FN={r["fn"]}, FP={r["fp"]}. '
            f"Validation AUC tại epoch loss tốt nhất: giai đoạn classifier {head.val_auc:.4f}; fine-tuning {fine.val_auc:.4f}.",
            f"![Loss](plots/{name}_loss_curve.png)",
            f"![Accuracy](plots/{name}_accuracy_curve.png)",
            f"![Precision Recall](plots/{name}_precision_recall.png)",
            f"![F1](plots/{name}_f1_curve.png)",
            f"![AUC](plots/{name}_auc_curve.png)",
            f"![Confusion matrix](confusion_matrix/{name}_confusion_matrix.png)",
        ]
        cases = json.loads((results / "gradcam" / f"{name}_cases.json").read_text())
        for case, count in cases.items():
            lines.append(
                f"{case}: {count} ảnh test."
                + (" Không có mẫu thuộc nhóm này." if count == 0 else "")
            )
            if count:
                lines.append(f"![{case}](gradcam/{name}/{case}_001_overlay.png)")
    lines += [
        "",
        "## Thảo luận và giới hạn",
        "Sensitivity cao hơn tương ứng tỷ lệ bỏ sót thấp hơn; cần xem đồng thời specificity và FP. "
        "So sánh trên một seed chưa chứng minh khác biệt có ý nghĩa thống kê. "
        "Kiểm tra heatmap ở vùng phổi, chữ, marker, viền và thiết bị để tìm dấu hiệu shortcut; "
        "không tự động suy ra vị trí tổn thương khi chưa có nhãn định vị hoặc chuyên gia đọc ảnh. "
        "Grad-CAM không chứng minh suy luận giống bác sĩ và không thay thế chẩn đoán lâm sàng. "
        "Patient ID suy từ tên tệp chưa được xác minh bằng metadata lâm sàng; kiểm tra trùng pixel không phát hiện mọi ảnh gần trùng. "
        "Dữ liệu nhi khoa từ một nguồn không chứng minh khả năng tổng quát sang người lớn hay bệnh viện khác. "
        "Chưa hiệu chỉnh xác suất và chưa đánh giá ngoài miền. Chưa kết luận tác dụng riêng của weighted loss nếu chưa chạy ablation.",
        "",
        "## Nguồn",
        "Kermany, Zhang, Goldbaum (2018). https://doi.org/10.17632/rscbjbr9sj.2 (CC BY 4.0).",
        "ResNet: https://arxiv.org/abs/1512.03385 ; DenseNet: https://arxiv.org/abs/1608.06993 ; Grad-CAM: https://arxiv.org/abs/1610.02391",
    ]
    (results / "report.md").write_text("\n\n".join(lines), encoding="utf-8")
    deck = Presentation()

    def slide(title, text="", picture=None):
        page = deck.slides.add_slide(deck.slide_layouts[5])
        page.shapes.title.text = title
        if picture:
            page.shapes.add_picture(str(picture), Inches(0.6), Inches(1.4), height=Inches(5))
        else:
            page.shapes.add_textbox(
                Inches(0.6), Inches(1.5), Inches(8.8), Inches(5)
            ).text_frame.text = text

    slide(
        "Phân loại viêm phổi từ X-quang",
        "Transfer Learning: ResNet50 / DenseNet121\nGrad-CAM\nKết quả thực nghiệm do pipeline tạo",
    )
    slide("Dữ liệu và phương pháp", "\n".join(lines[5:6]))
    slide("Phân bố dữ liệu", picture=results / "dataset_distribution.png")
    slide("Kết quả test", picture=results / "plots" / "model_comparison.png")
    page = deck.slides.add_slide(deck.slide_layouts[5])
    page.shapes.title.text = "Bảng kết quả test thực nghiệm"
    table = page.shapes.add_table(3, 7, Inches(0.25), Inches(1.7), Inches(9.5), Inches(2)).table
    for column, title in enumerate(
        ["Model", "Accuracy", "Precision", "Sensitivity", "Specificity", "F1", "AUC"]
    ):
        table.cell(0, column).text = title
    for row, record in enumerate(records, 1):
        table.cell(row, 0).text = record["model"]
        for column, key in enumerate(fields[:6], 1):
            table.cell(row, column).text = f"{record[key]:.4f}"
    slide("ROC trên test", picture=results / "roc" / "model_comparison_roc.png")
    for name in names:
        for suffix in ("loss_curve", "accuracy_curve", "precision_recall", "f1_curve", "auc_curve"):
            slide(f"{name}: {suffix}", picture=results / "plots" / f"{name}_{suffix}.png")
        slide(
            f"{name}: Confusion matrix",
            picture=results / "confusion_matrix" / f"{name}_confusion_matrix.png",
        )
        for case in ("TP", "TN", "FP", "FN"):
            path = results / "gradcam" / name / f"{case}_001_overlay.png"
            if path.exists():
                slide(f"{name}: {case} Grad-CAM", picture=path)
    slide(
        "Kết luận và giới hạn",
        f'Chọn bằng validation: {best["model"]}\nXem đồng thời AUC, Sensitivity, Specificity, F1, FN.\nMột seed; chưa external validation hoặc calibration.\nGrad-CAM không thay thế chẩn đoán lâm sàng.\nNguồn: Kermany et al., doi:10.17632/rscbjbr9sj.2',
    )
    deck.save(results / "presentation.pptx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default=".")
    report(parser.parse_args().output)
